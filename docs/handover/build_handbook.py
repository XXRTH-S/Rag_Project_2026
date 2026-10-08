"""Build the Thai handover document from the adjacent Markdown and config AST."""
from pathlib import Path
import ast
import json
import re
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from PIL import Image, ImageDraw, ImageFont

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent.parent
QA = HERE / '_qa'
QA.mkdir(exist_ok=True)
FONT = 'Leelawadee UI'
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.top_margin, sec.bottom_margin = Cm(2), Cm(1.8)
sec.left_margin = sec.right_margin = Cm(1.8)
sec.footer_distance = Cm(0.8)

def set_font(style, size, bold=False):
    style.font.name = FONT
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
        fonts.set(qn('w:'+key), FONT)
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'th-TH'); rpr.append(lang)
    cs = OxmlElement('w:szCs'); cs.set(qn('w:val'), str(int(size*2))); rpr.append(cs)

set_font(doc.styles['Normal'], 10.5)
normal = doc.styles['Normal'].paragraph_format
normal.line_spacing = 1.2
normal.space_after = Pt(7)
normal.widow_control = True
for name, size in [('Title',26),('Subtitle',14),('Heading 1',18),('Heading 2',13)]:
    set_font(doc.styles[name],size,name != 'Subtitle')
    doc.styles[name].paragraph_format.space_before = Pt(12)
    doc.styles[name].paragraph_format.space_after = Pt(8)
    doc.styles[name].paragraph_format.keep_with_next = True
doc.styles['Heading 1'].paragraph_format.page_break_before = True
for name in ['List Bullet','List Number']:
    set_font(doc.styles[name],10.5)
    doc.styles[name].paragraph_format.line_spacing = 1.2
    doc.styles[name].paragraph_format.space_after = Pt(5)
code_style = doc.styles.add_style('Code Block', 1)
set_font(code_style,9)
code_style.font.name = 'Consolas'
code_style.paragraph_format.space_after = Pt(1)
code_style.paragraph_format.line_spacing = 1.1
set_font(doc.styles['Caption'],9)
doc.styles['Caption'].font.italic=False

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r=footer.add_run('RAG Workshop  |  '); r.font.size=Pt(8)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='คู่มือส่งมอบระบบ RAG Workshop'
doc.core_properties.subject='Architecture Flow Configuration and Model Tuning'
doc.core_properties.author=''
doc.core_properties.keywords='RAG, Junior Developer, Thai, Handover'

def shade(el, color):
    s=OxmlElement('w:shd'); s.set(qn('w:fill'),color); el.append(s)

def table(rows, widths=None):
    n=len(rows[0]); t=doc.add_table(rows=1,cols=n)
    t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    widths=widths or ([5.3,5.5,6.6] if n==3 else [8.7,8.7])
    for c,w in zip(t.columns,widths): c.width=Cm(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for side in ['top','left','bottom','right','insideH','insideV']:
        b=OxmlElement('w:'+side); b.set(qn('w:val'),'single'); b.set(qn('w:sz'),'4'); b.set(qn('w:color'),'D9D9D9'); borders.append(b)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for side,val in [('top',85),('bottom',85),('left',95),('right',95)]:
        m=OxmlElement('w:'+side); m.set(qn('w:w'),str(val)); m.set(qn('w:type'),'dxa'); margins.append(m)
    pr.append(margins)
    for idx,values in enumerate(rows):
        row=t.rows[0] if idx==0 else t.add_row()
        trpr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit'); trpr.append(cant)
        if idx==0:
            repeat=OxmlElement('w:tblHeader'); trpr.append(repeat)
        for cell,txt,w in zip(row.cells,values,widths):
            cell.width=Cm(w); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell.text=txt
            shade(cell._tc.get_or_add_tcPr(),'19394A' if idx==0 else ('F2F5F7' if idx%2==0 else 'FFFFFF'))
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.15
                for r in p.runs:
                    r.font.size=Pt(9)
                    if idx==0: r.font.bold=True; r.font.color.rgb=RGBColor(255,255,255)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

notes={
'model_tier':'ป้ายชื่อ tier ไม่ได้เลือก compose',
'llm_base_url':'URL ของ chat provider', 'llm_api_key':'credential ของ provider',
'llm_model':'model ID สำหรับ chat','llm_max_model_len':'ประกาศไว้ แต่ client ไม่ส่งปรับ runtime',
'llm_temperature':'การสุ่ม output','llm_max_tokens':'จำนวน output tokens สูงสุด',
'llm_timeout_seconds':'HTTP timeout หน่วยวินาที','llm_enable_thinking':'การเปิด thinking ใน client payload',
'llm_reasoning_effort':'field ตามความสามารถ provider',
'llm_fallback_base_url':'URL สำรอง ว่างคือปิด','llm_fallback_api_key':'credential สำรอง','llm_fallback_model':'model สำรอง ว่างใช้ชื่อหลัก',
'llm_breaker_threshold':'จำนวน HTTP failures ก่อนเปิด breaker','llm_breaker_cooldown_seconds':'ช่วงพัก breaker วินาที',
'llm_fallback_to_excerpts':'ยกข้อความแทนเมื่อ LLM ใช้ไม่ได้','llm_warmup_on_start':'warmup เมื่อ startup เปิด ingestion',
'typhoon_ocr_base_url':'URL ของ OCR provider','typhoon_ocr_api_key':'credential OCR','typhoon_ocr_model':'model ID หรือ alias OCR',
'typhoon_ocr_task_type':'default หรือ structure','typhoon_ocr_concurrency':'ประกาศไว้ ยังไม่คุม loop งานจริง','typhoon_ocr_page_timeout':'HTTP timeout ของ OCR วินาที',
'ocr_min_chars_per_page':'เกณฑ์อักขระก่อนเลือก OCR','ocr_seconds_per_page':'ค่าประมาณ ETA ต่อหน้า ไม่ใช่ timeout',
'ingestion_stale_after_minutes':'อายุ queued ที่ยังไม่ started สำหรับ reprocess',
'embedding_base_url':'URL ของ TEI /embed','embedding_model':'ชื่อแสดงผล รุ่นจริงตั้งใน TEI','embedding_dim':'มิติที่ client ตรวจ ต้องตรง schema',
'embedding_device':'ป้าย config ไม่ได้เปลี่ยนอุปกรณ์ TEI','embedding_api_key':'credential embedding endpoint',
'database_url':'DB connection ต้องตั้ง credential จริง','redis_url':'broker backend และ rate counters','db_pool_enabled':'ใช้ pool หรือ NullPool',
'quota_timezone':'วันรอบโควตา','user_daily_document_limit':'เอกสารต่อวัน','user_daily_page_limit':'หน้าต่อวัน',
'user_max_pages_per_document':'หน้าสูงสุดต่อไฟล์สำหรับ user','admin_unlimited':'ข้ามการบล็อก quota ของ admin',
'chars_per_page_estimate':'ตัวอักษรสำหรับประมาณหนึ่งหน้า',
'retrieval_top_k':'จำนวน hits สูงสุด','retrieval_min_score':'cosine similarity floor',
'retrieval_hybrid':'เปิด vector ร่วม keyword','retrieval_keyword_floor':'floor สำหรับ keyword match',
'chunk_size':'เป้าหมายตัวอักษรต่อ chunk','chunk_overlap':'ตัวอักษร overlap',
'ingestion_enabled':'เปิด upload bulk reprocess','app_secret_key':'JWT signing key ต้องตั้งใหม่',
'admin_email':'bootstrap admin account','admin_password':'bootstrap และ sync password admin',
'cors_allowed_origins':'origin คั่นด้วย comma','public_api_url':'URL public และ compose frontend mapping',
'upload_dir':'persistent path ใน container','max_upload_mb':'เพดาน MB ต่อไฟล์','access_token_expire_minutes':'อายุ JWT หน่วยนาที',
'cookie_secure':'จำกัด cookie ผ่าน HTTPS','api_docs_enabled':'เปิด docs OpenAPI และ ReDoc',
'bot_polite_particle':'คำลงท้ายภาษาไทย','widget_title':'หัว widget','widget_greeting':'คำทักทาย widget',
'widget_accent_color':'สี widget','widget_suggestions':'คำถามตัวอย่าง คั่นด้วยเครื่องหมาย pipe',
'rate_limit_chat_per_minute':'chat ต่อ user ต่อนาที','rate_limit_upload_per_hour':'upload ต่อ user ต่อชั่วโมง',
'rate_limit_login_attempts':'login ต่อ IP ต่อ window','rate_limit_login_window_seconds':'window login ต่อ IP วินาที',
'rate_limit_login_attempts_per_email':'login ต่ออีเมล 0 คือปิดชั้นนี้','rate_limit_login_email_window_seconds':'window ต่ออีเมล วินาที',
'rate_limit_playground_per_minute':'Playground ต่อ user ต่อนาที','trusted_proxy_hops':'จำนวน proxy ที่เชื่อถือ นับจากขวา',
}
config=ast.parse((ROOT/'api/app/core/config.py').read_text(encoding='utf-8-sig'))
fields=[]
for cls in config.body:
    if isinstance(cls,ast.ClassDef) and cls.name=='Settings':
        for item in cls.body:
            if isinstance(item,ast.AnnAssign) and isinstance(item.target,ast.Name):
                key=item.target.id
                if isinstance(item.value,ast.BinOp):
                    value=ast.literal_eval(item.value.left)*ast.literal_eval(item.value.right)
                else: value=ast.literal_eval(item.value)
                if key in ('app_secret_key','database_url'): value='ต้องกำหนดค่าปลอดภัย'
                elif isinstance(value,bool): value=str(value).lower()
                elif value=='': value='ว่าง'
                fields.append((key.upper(),str(value),notes[key]))

def make_diagram(kind):
    im=Image.new('RGB',(1800,1000),'white'); d=ImageDraw.Draw(im)
    f=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',29)
    fb=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',32)
    def box(rect,title,sub=''):
        d.rounded_rectangle(rect,18,fill='#eef3f6',outline='#244858',width=3)
        x1,y1,x2,y2=rect
        lines=[title]+sub.split('\n') if sub else [title]
        for j,line in enumerate(lines):
            font=fb if j==0 else f
            bbox=d.textbbox((0,0),line,font=font)
            d.text(((x1+x2-(bbox[2]-bbox[0]))/2,y1+20+j*42),line,font=font,fill='#12242e')
    def arrow(points,label='',lx=None,ly=None):
        d.line(points,fill='#496979',width=4)
        x,y=points[-1]; px,py=points[-2]
        if x==px:
            k=1 if y>py else -1; tri=[(x,y),(x-9,y-k*16),(x+9,y-k*16)]
        else:
            k=1 if x>px else -1; tri=[(x,y),(x-k*16,y-9),(x-k*16,y+9)]
        d.polygon(tri,fill='#496979')
        if label:d.text((lx,ly),label,font=f,fill='#244858')
    if kind==0:
        box((40,70,360,185),'Browser / Widget')
        box((540,70,880,185),'Caddy',':80  reverse proxy')
        box((1140,70,1720,185),'Next.js',':3000  user interface')
        box((540,320,880,445),'FastAPI',':8000  auth + RAG')
        box((40,580,360,705),'Redis','queues + counters')
        box((40,800,360,925),'Celery worker','concurrency = 1')
        box((580,800,900,925),'Uploads','original files')
        box((1170,310,1720,435),'PostgreSQL + pgvector','metadata + vectors + chat')
        box((1170,555,1720,680),'TEI on CPU','bge-m3  /embed')
        box((1170,800,1720,925),'Ollama on GPU','chat model + OCR model')
        arrow([(360,127),(540,127)])
        arrow([(880,127),(1140,127)],'pages',940,80)
        arrow([(710,185),(710,320)],'/api/*',730,235)
        arrow([(880,380),(1170,380)])
        arrow([(880,415),(1020,415),(1020,615),(1170,615)])
        arrow([(880,430),(980,430),(980,850),(1170,850)])
        arrow([(540,410),(200,410),(200,580)],'enqueue',215,465)
        arrow([(200,705),(200,800)])
        arrow([(360,865),(580,865)])
        d.text((400,950),'Worker also calls TEI, Ollama and PostgreSQL',font=f,fill='#244858')
    elif kind==1:
        box((50,80,460,210),'users','identity + ownership')
        box((690,80,1110,210),'documents','status + storage_path')
        box((1350,80,1760,210),'chunks','text + embedding')
        box((690,390,1110,520),'ingestion_jobs','stage + progress')
        box((50,390,460,520),'quota tables','counter + event trail')
        box((50,720,460,850),'chat_sessions','owner + title')
        box((690,720,1110,850),'chat_messages','content + model + timing')
        box((1350,580,1760,710),'message_citations','source metadata')
        box((1350,820,1760,950),'feedback','rating + comment')
        arrow([(460,145),(690,145)],'1 : N',540,100)
        arrow([(1110,145),(1350,145)],'1 : N',1175,100)
        arrow([(900,210),(900,390)],'1 : N',920,290)
        arrow([(250,210),(250,390)],'1 : N',270,290)
        arrow([(85,210),(20,210),(20,785),(50,785)])
        arrow([(460,785),(690,785)],'1 : N',540,735)
        arrow([(1110,760),(1240,760),(1240,650),(1350,650)])
        arrow([(1110,810),(1230,810),(1230,885),(1350,885)])
        d.text((535,950),'Prompt configs are referenced by chat messages.',font=f,fill='#244858')
    else:
        box((60,80,480,210),'Upload accepted','HTTP 202 + job ID')
        box((700,80,1120,210),'Queue','commit before enqueue')
        box((1340,80,1760,210),'Detect','inspect each PDF page')
        box((1340,390,1760,520),'Parse or OCR','read text / image to text')
        box((700,390,1120,520),'Clean and chunk','text + page + source')
        box((60,390,480,520),'Embed and store','TEI + PostgreSQL')
        box((60,750,480,880),'Ready / Done','eligible for retrieval')
        box((700,750,1120,880),'Failed','log reference + inspection')
        arrow([(480,145),(700,145)])
        arrow([(1120,145),(1340,145)])
        arrow([(1550,210),(1550,390)])
        arrow([(1340,455),(1120,455)])
        arrow([(700,455),(480,455)])
        arrow([(270,520),(270,750)])
        arrow([(900,520),(900,750)],'on error',925,620)
        d.text((1170,790),'Actual UI stages are approximate;',font=f,fill='#244858')
        d.text((1170,835),'see the pipeline notes in Chapter 15.',font=f,fill='#244858')
    path=QA/f'diagram-{kind+1}.png'; im.save(path,dpi=(240,240));return path

md=(HERE/'RAG_Workshop_Junior_Handbook_TH.md').read_text(encoding='utf-8')
config_md='| Environment | Default ใน source | ความหมาย |\n| --- | --- | --- |\n'+'\n'.join('| '+' | '.join(c.replace('|',' / ') for c in row)+' |' for row in fields)
if '{{CONFIG_TABLE}}' in md:
    md=md.replace('{{CONFIG_TABLE}}',config_md)
else:
    md=re.sub(r'\| Environment \| Default ใน source \| ความหมาย \|.*?(?=\n\n## 19 )',lambda _:config_md,md,flags=re.S)
(HERE/'RAG_Workshop_Junior_Handbook_TH.md').write_text(md,encoding='utf-8')
lines=md.splitlines(); i=0; diagram_no=0; in_toc=False; bookmark_id=0
while i<len(lines):
    line=lines[i].strip()
    if not line: i+=1;continue
    if line.startswith('```'):
        kind=line[3:]; block=[]; i+=1
        while i<len(lines) and not lines[i].startswith('```'):
            block.append(lines[i]);i+=1
        if kind=='diagram':
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            pic=p.add_run().add_picture(str(make_diagram(diagram_no)),width=Cm(17.3))
            pic._inline.docPr.set('descr',['Architecture services and data flow','Entity relationships','Document ingestion lifecycle'][diagram_no])
            doc.add_paragraph(['ภาพ 1 องค์ประกอบระบบและบริการที่ติดต่อกัน','ภาพ 2 ความสัมพันธ์ของข้อมูลหลัก','ภาพ 3 ลำดับการประมวลผลเอกสาร'][diagram_no],'Caption')
            diagram_no+=1
        else:
            for n,text in enumerate(block):
                p=doc.add_paragraph(text,'Code Block')
                p.paragraph_format.keep_with_next=n<len(block)-1
                shade(p._p.get_or_add_pPr(),'F2F5F7')
        i+=1;continue
    if line.startswith('|'):
        rows=[]
        while i<len(lines) and lines[i].strip().startswith('|'):
            vals=[v.strip() for v in lines[i].strip().strip('|').split('|')]
            if not all(re.match(r'^[-: ]+$',v) for v in vals):rows.append(vals)
            i+=1
        widths=[6.2,4.1,7.1] if rows[0][0]=='Environment' else None
        table(rows,widths);continue
    if line.startswith('# '):
        doc.add_paragraph(line[2:],'Title')
    elif line.startswith('## '):
        title=line[3:]; p=doc.add_paragraph(title,'Heading 1')
        in_toc=title=='สารบัญและเส้นทางการอ่าน'
        num=re.match(r'^(\d+) ',title)
        if num:
            bookmark_id+=1
            start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(bookmark_id));start.set(qn('w:name'),'chapter_'+num[1]);p._p.insert(0,start)
            end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bookmark_id));p._p.append(end)
    elif line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2')
    elif line.startswith('- '):doc.add_paragraph(line[2:],'List Bullet')
    elif in_toc and re.match(r'^\d+\. ',line):
        p=doc.add_paragraph();p.paragraph_format.space_after=Pt(3)
        link=OxmlElement('w:hyperlink');link.set(qn('w:anchor'),'chapter_'+line.split('.')[0])
        run=OxmlElement('w:r');txt=OxmlElement('w:t');txt.text=line;run.append(txt);link.append(run);p._p.append(link)
    elif re.match(r'^\d+\. ',line):doc.add_paragraph(line,'Normal')
    else:doc.add_paragraph(line)
    i+=1

out=HERE/'RAG_Workshop_Junior_Handbook_TH.docx'
doc.save(out)
check=Document(out)
report={'config_fields':len(fields),'paragraphs':len(check.paragraphs),'tables':len(check.tables),'figures':len(check.inline_shapes),'chapters':sum(p.style.name=='Heading 1' for p in check.paragraphs)-1,'docx_bytes':out.stat().st_size,'render_verified':False}
assert report['chapters']==20
assert report['figures']==3
assert all(len(r.cells)==len(t.columns) for t in check.tables for r in t.rows)
assert '{{CONFIG_TABLE}}' not in md
(QA/'structure-check.json').write_text(json.dumps(report,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps(report,ensure_ascii=False))
